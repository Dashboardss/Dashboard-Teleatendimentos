from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, date, time
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path('/home/ubuntu/upload/ESCALATELEAUTÔNOMO2026(2).xlsx')
OUT = Path('/home/ubuntu/dashboard_teleautonomo_2026')
OUT.mkdir(parents=True, exist_ok=True)

MONTHS = {
    'JANEIRO': 1, 'FEVEREIRO': 2, 'MARÇO': 3, 'MARCO': 3,
    'ABRIL': 4, 'MAIO': 5, 'JUNHO': 6, 'JULHO': 7,
    'AGOSTO': 8, 'SETEMBRO': 9, 'OUTUBRO': 10, 'NOVEMBRO': 11,
    'DEZEMBRO': 12,
}


def norm(v):
    if v is None:
        return ''
    s = str(v).strip()
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'\s+', ' ', s).strip().lower()


def clean_text(v):
    if v is None:
        return ''
    return re.sub(r'\s+', ' ', str(v).strip())


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = clean_text(v)
    m = re.search(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def parse_time(v):
    if isinstance(v, datetime):
        return v.strftime('%H:%M')
    if isinstance(v, time):
        return v.strftime('%H:%M')
    s = clean_text(v)
    m = re.search(r'(\d{1,2}):(\d{2})', s)
    return f'{int(m.group(1)):02d}:{m.group(2)}' if m else s


def status_label(v, name=''):
    raw = clean_text(v)
    n = norm(raw)
    nn = norm(name)
    combined = f'{n} {nn}'
    if 'vaga' in combined and 'livre' in combined:
        return 'Vagas livres'
    if 'nao atend' in combined or 'nao atendi' in combined:
        return 'Não atendido'
    if 'desist' in combined:
        return 'Desistência'
    if 'cancel' in combined:
        return 'Cancelado'
    if re.search(r'\batendid', combined):
        return 'Atendido'
    if 'remanejad' in combined:
        return 'Remanejado'
    if n in ('feriado',):
        return 'Feriado'
    if raw:
        return raw
    if name and nn not in ('', 'vagas livres'):
        return 'Sem status'
    return ''


def professional_from_sheet(sheet_name):
    n = norm(sheet_name)
    if 'camila' in n:
        return 'Dra. Camila Queiroga'
    if 'viridiana' in n:
        return 'Dras. Viridiana & Priscila'
    if 'priscila' in n:
        return 'Dra. Priscila Ferreira'
    return 'Não identificado'


def sheet_month(sheet_name):
    n = norm(sheet_name).upper()
    for name, number in MONTHS.items():
        if name in n:
            return number
    return None


def is_operational_2026(name):
    n = norm(name)
    return '26' in n and any(x in n for x in ('camila', 'priscila', 'viridiana'))

wb = load_workbook(SOURCE, data_only=True, read_only=False)
records = []
sheet_audit = []

for sheet_name in wb.sheetnames:
    if not is_operational_2026(sheet_name):
        continue
    ws = wb[sheet_name]
    professional = professional_from_sheet(sheet_name)
    smonth = sheet_month(sheet_name)
    current_date = None
    header = None
    header_row = None
    sheet_records = 0
    status_counts = Counter()

    for row in ws.iter_rows():
        vals = [c.value for c in row]
        vals = vals + [None] * max(0, 18 - len(vals))
        first = clean_text(vals[0])
        found_date = parse_date(first)
        is_date_header = (bool(re.search(r'\d{1,2}/\d{1,2}/\d{4}', first)) and ('h' in first.lower() or 'ás' in first.lower() or 'as' in first.lower())) or (isinstance(vals[0], (datetime, date)) and not any(clean_text(v) for v in vals[1:]))
        # Excel may expose numeric slot IDs as datetimes when the cell has an
        # accidental date format. Treat a datetime as a section date only when
        # the rest of the row is empty (the actual date header pattern).
        if isinstance(vals[0], (datetime, date)) and any(clean_text(v) for v in vals[1:]):
            is_date_header = False
        if found_date and is_date_header:
            current_date = found_date if found_date.year == 2026 else None
            header = None
            header_row = None
            continue

        norm_vals = [norm(v) for v in vals]
        if any(v == 'status' or v.startswith('status ') for v in norm_vals):
            header = {v.rstrip(): i for i, v in enumerate(norm_vals) if v}
            header_row = row[0].row
            continue
        if header is None or current_date is None or row[0].row <= header_row:
            continue

        status_idx = next((i for k, i in header.items() if k == 'status' or k.startswith('status')), 8)
        unit_idx = next((i for k, i in header.items() if 'unidade' in k), 1)
        time_idx = next((i for k, i in header.items() if 'horario' in k or 'horário' in k), 3)
        name_idx = next((i for k, i in header.items() if 'nome completo' in k), 6)
        status_raw = vals[status_idx] if status_idx < len(vals) else None
        name = clean_text(vals[name_idx] if name_idx < len(vals) else None)
        unit = clean_text(vals[unit_idx] if unit_idx < len(vals) else None)
        slot = clean_text(vals[0])
        hour = parse_time(vals[time_idx] if time_idx < len(vals) else None)

        # Ignore separator/summary rows and preserve actual slots, including vacancies.
        has_slot = bool(re.fullmatch(r'\d+(?:\.0)?', slot))
        if not has_slot:
            continue
        if not any(clean_text(v) for v in vals[1:12]) and not status_raw:
            continue
        status = status_label(status_raw, name)
        if not status:
            continue
        # Do not include raw patient identifiers; only operational dimensions remain.
        rec = {
            'date': current_date.isoformat(),
            'sheet_month': smonth,
            'month': current_date.month,
            'month_label': current_date.strftime('%m/%Y'),
            'professional': professional,
            'unit': unit or 'Não informado',
            'weekday': current_date.strftime('%A'),
            'hour': hour or 'Não informado',
            'status': status,
        }
        records.append(rec)
        sheet_records += 1
        status_counts[status] += 1

    sheet_audit.append({
        'sheet': sheet_name,
        'professional': professional,
        'sheet_month': smonth,
        'records': sheet_records,
        'statuses': dict(status_counts),
    })

# Translate weekdays to PT-BR for the public artifact.
weekday_pt = {
    'Monday': 'Segunda-feira', 'Tuesday': 'Terça-feira', 'Wednesday': 'Quarta-feira',
    'Thursday': 'Quinta-feira', 'Friday': 'Sexta-feira', 'Saturday': 'Sábado', 'Sunday': 'Domingo'
}
for r in records:
    r['weekday'] = weekday_pt.get(r['weekday'], r['weekday'])

status_order = ['Atendido', 'Não atendido', 'Cancelado', 'Desistência', 'Vagas livres', 'Remanejado', 'Feriado', 'Sem status']

def counts(key):
    c = Counter(r[key] for r in records)
    return [{'label': k, 'value': c[k]} for k in status_order if c[k]] if key == 'status' else [{'label': k, 'value': v} for k, v in sorted(c.items())]


def aggregate(keys):
    c = Counter(tuple(r[k] for k in keys) for r in records)
    return [dict(zip(keys, key), **{'value': value}) for key, value in sorted(c.items())]

attendance = [r for r in records if r['status'] in ('Atendido', 'Não atendido', 'Cancelado', 'Desistência')]
realized = sum(r['status'] == 'Atendido' for r in records)
not_attended = sum(r['status'] == 'Não atendido' for r in records)
cancelled = sum(r['status'] in ('Cancelado', 'Desistência') for r in records)
slots = len([r for r in records if r['status'] not in ('Vagas livres', 'Remanejado', 'Feriado')])

payload = {
    'meta': {
        'source_file': SOURCE.name,
        'generated_at': datetime.now().isoformat(timespec='seconds'),
        'scope': 'Abas operacionais com referência a 2026 e profissional identificado',
        'privacy': 'Dashboard agregado; não contém nome, CPF, telefone, cartão SUS ou data de nascimento.',
        'sheets_analyzed': len(sheet_audit),
        'records': len(records),
    },
    'records': records,
    'kpis': {
        'records': len(records),
        'scheduled': slots,
        'realized': realized,
        'not_attended': not_attended,
        'cancelled_or_withdrawn': cancelled,
        'vacancies': sum(r['status'] == 'Vagas livres' for r in records),
        'attendance_rate': round(realized / slots * 100, 1) if slots else 0,
        'loss_rate': round((not_attended + cancelled) / slots * 100, 1) if slots else 0,
    },
    'status': counts('status'),
    'by_month': aggregate(['month_label', 'status']),
    'by_date': aggregate(['date', 'status']),
    'by_professional': aggregate(['professional', 'status']),
    'by_unit': aggregate(['unit', 'status']),
    'by_hour': aggregate(['hour', 'status']),
    'audit': sheet_audit,
}
(OUT / 'data.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')

# Compact CSV for downstream use, without any personal data.
import csv
with (OUT / 'atendimentos_agregados.csv').open('w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=['date','month_label','professional','unit','weekday','hour','status'])
    w.writeheader()
    w.writerows({k: r.get(k, '') for k in w.fieldnames} for r in records)

print(json.dumps({
    'sheets_analyzed': len(sheet_audit),
    'records': len(records),
    'kpis': payload['kpis'],
    'statuses': payload['status'],
    'professionals': sorted(set(r['professional'] for r in records)),
}, ensure_ascii=False, indent=2))
